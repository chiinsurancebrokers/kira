"""
dossier.py — "Φάκελος για τον γιατρό" (doctor dossier) for Asklepios.

Turns a pile of exam files (PDFs, phone photos of printouts) into one clean,
doctor-facing summary: patient header, current medication table, one section
per exam (date · centre · clinician, parameter/value tables, the recorded
conclusion exactly as written) and a short list of "points that may be worth
mentioning" — factual co-occurrences only, never interpretation or advice.

Output formats: Word (.docx, python-docx) and a self-contained printable HTML
page (the user prints it to PDF from the phone/browser).

Design rules that make doctors trust it:
  * transcribe, don't interpret — values, units and wording are copied as
    printed; the model is forbidden from inventing, rounding or "fixing";
  * every exam carries its source (date, centre, clinician, report status);
  * anything illegible is marked, never guessed;
  * the only synthesis is the "points worth mentioning" list, which is framed
    as observations for completeness, with no diagnosis or recommendation.

Privacy: files are sent to the Claude API for extraction and kept only in the
user's Streamlit session — nothing is written to disk or a database here.
"""
from __future__ import annotations

import base64
import concurrent.futures as _cf
import html as _html
import io
import json
import re
import time
import urllib.request
from datetime import datetime

MODEL = "claude-sonnet-4-6"
API_URL = "https://api.anthropic.com/v1/messages"

# ── Labels ────────────────────────────────────────────────────────────────────
L = {
    "el": {
        "title": "Συγκεντρωτικά Στοιχεία Εξετάσεων",
        "item": "Στοιχείο", "value": "Τιμή",
        "patient": "Ασθενής", "amka": "ΑΜΚΑ", "agesex": "Ηλικία / Φύλο",
        "history": "Γνωστό ιστορικό", "dx": "Διάγνωση / λόγος επίσκεψης",
        "years": "ετών",
        "per_exam": "Σύνοψη ακολουθεί ανά εξέταση, με ημερομηνία και κέντρο διενέργειας.",
        "meds_h": "Τρέχουσα Φαρμακευτική Αγωγή",
        "meds_note": "Όπως καταγράφηκε από τον ασθενή / την οικογένεια",
        "time": "Ώρα", "drug": "Φάρμακο", "dose": "Δοσολογία",
        "symptoms_h": "Πρόσφατα Συμπτώματα",
        "symptoms_note": "Όπως τα περιέγραψε ο ασθενής στη συνομιλία με τον Asklepios — χωρίς ερμηνεία.",
        "param": "Παράμετρος",
        "conclusion": "Συμπέρασμα εξέτασης (όπως καταγράφηκε)",
        "finding": "Εύρημα", "details": "Λεπτομέρειες",
        "source_note": "Σημείωση πηγής",
        "points_h": "Σημεία που Ίσως Αξίζει να Αναφερθούν",
        "points_note": ("Παρατηρήσεις από τη συνδυαστική εξέταση των παραπάνω στοιχείων — "
                        "όχι ερμηνεία, απλώς προς επισήμανση εφόσον είναι χρήσιμο."),
        "footer": ("Τα παραπάνω στοιχεία συγκεντρώθηκαν από τα πρωτότυπα αρχεία εξετάσεων για "
                   "διευκόλυνση της επισκόπησης πριν το ραντεβού. Δεν περιλαμβάνουν ερμηνεία ή "
                   "σύσταση — τα πρωτότυπα αρχεία παραμένουν διαθέσιμα εφόσον χρειαστούν."),
        "made_with": "Δημιουργήθηκε με τον Asklepios · AI Νοσηλευτή",
        "prepared": "Ημερομηνία σύνταξης",
        "trend_h": "Εξέλιξη Τιμών",
        "trend_note": "Οι ίδιες παράμετροι σε διαφορετικές ημερομηνίες, όπως είναι γραμμένες στις εξετάσεις (νεότερη πρώτη).",
    },
    "en": {
        "title": "Summary of Examination Results",
        "item": "Item", "value": "Value",
        "patient": "Patient", "amka": "Social security no. (AMKA)", "agesex": "Age / Sex",
        "history": "Known history", "dx": "Diagnosis / reason for visit",
        "years": "years",
        "per_exam": "A summary follows per examination, with date and centre.",
        "meds_h": "Current Medication",
        "meds_note": "As reported by the patient / family",
        "time": "Time", "drug": "Medication", "dose": "Dose",
        "symptoms_h": "Recent Symptoms",
        "symptoms_note": "As described by the patient in the conversation with Asklepios — no interpretation.",
        "param": "Parameter",
        "conclusion": "Examination conclusion (as recorded)",
        "finding": "Finding", "details": "Details",
        "source_note": "Source note",
        "points_h": "Points That May Be Worth Mentioning",
        "points_note": ("Observations from looking at the results above together — not an "
                        "interpretation, only flagged in case it is useful."),
        "footer": ("The information above was compiled from the original examination files to make "
                   "review easier before the appointment. It contains no interpretation or "
                   "recommendation — the original files remain available if needed."),
        "made_with": "Prepared with Asklepios · AI Nurse",
        "prepared": "Prepared on",
        "trend_h": "Values Over Time",
        "trend_note": "The same parameters on different dates, as printed in the exams (newest first).",
    },
}


def labels(lang: str) -> dict:
    return L["el"] if lang == "el" else L["en"]


# ── Claude call ───────────────────────────────────────────────────────────────
def _call_claude(api_key: str, system: str, content: list, max_tokens: int = 8000,
                 timeout: int = 150) -> str:
    body = json.dumps({
        "model": MODEL, "max_tokens": max_tokens, "system": system,
        "messages": [{"role": "user", "content": content}],
    }).encode()
    req = urllib.request.Request(API_URL, data=body, headers={
        "x-api-key": api_key, "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    })
    with urllib.request.urlopen(req, timeout=timeout) as r:
        data = json.loads(r.read())
    return "".join(b.get("text", "") for b in data.get("content", []) if b.get("type") == "text")


def _parse_json(text: str):
    """Pull the first JSON object out of a model reply (tolerates code fences)."""
    if not text:
        raise ValueError("empty reply")
    t = text.strip()
    t = re.sub(r"^```(?:json)?\s*", "", t)
    t = re.sub(r"\s*```$", "", t)
    try:
        return json.loads(t)
    except Exception:
        a, b = t.find("{"), t.rfind("}")
        if a >= 0 and b > a:
            return json.loads(t[a:b + 1])
        raise


# ── 1. Extract exams from one file ────────────────────────────────────────────
_EXTRACT_SYSTEM = """You are a meticulous medical records transcriptionist. You convert examination
reports (PDF or photos) into structured JSON so they can be compiled into a summary for the
patient's doctor. You TRANSCRIBE — you never interpret, diagnose, recommend, round, correct or
complete values. Output JSON only."""

def _extract_prompt(lang: str) -> str:
    target = "Greek" if lang == "el" else "English"
    return f"""Extract EVERY examination contained in the attached document into this JSON:

{{
  "patient_name_on_document": string|null,
  "exams": [
    {{
      "exam_title": string,            // e.g. "Ηχοκαρδιογράφημα (Διαθωρακικό)", "Γενική αίματος"
      "exam_type": string,             // short english category: imaging|echo|ecg|holter|stress|angiography|labs|pathology|endoscopy|other
      "date": string|null,             // DD/MM/YYYY exactly as the document dates the exam
      "date_iso": string|null,         // YYYY-MM-DD, null if unknown
      "facility": string|null,         // hospital / lab / diagnostic centre
      "clinician": string|null,        // performing / signing doctor with specialty, as printed
      "source_note": string|null,      // report status as printed, e.g. «Unconfirmed Report», «Προσωρινό»
      "sections": [
        {{
          "title": string|null,        // sub-heading as in the report, e.g. "Doppler — Αορτική βαλβίδα"
          "description": string|null,  // free-text description under that heading, verbatim
          "columns": [string, ...],    // 2 or 3 headers, e.g. ["Παράμετρος","Τιμή"] or ["Παράμετρος","Τιμή","Φυσιολογικά όρια"]
          "rows": [[string, ...], ...] // one row per measured item; same length as columns
        }}
      ],
      "conclusion": [                  // the report's own conclusion / impression, VERBATIM meaning
        {{"item": string, "detail": string}}
      ]
    }}
  ]
}}

Rules:
- Write labels, headings and descriptions in {target}. Keep medical abbreviations the report uses
  (EF, AVA, LVDD, PR, QTc…) in parentheses after the {target} label.
- Copy every number, unit, range, date and sign EXACTLY as printed. Never round, convert or infer.
- Include reference ranges as a third column when the report prints them; mark out-of-range
  values only if the report itself flags them (e.g. "H", "*", "↑") — copy that flag.
- Group rows under the report's own sub-headings. Lab panels: one section per panel.
- If something is illegible write "[δυσανάγνωστο]" (Greek) / "[illegible]" (English). Never guess.
- If the document has several distinct examinations (e.g. ECG + stress echo), return several exams.
- If the page is not a medical examination (e.g. an invoice), return "exams": [].
- Do NOT add interpretation, advice or your own conclusions. "conclusion" only holds what the
  report itself concludes; leave it [] if it has none.
Return JSON only."""


def _prep_file(file_bytes: bytes, mime: str, name: str, heic_convert=None, downscale=None):
    """Return a Claude content block for a PDF or image."""
    low = (name or "").lower()
    if mime == "application/pdf" or low.endswith(".pdf"):
        return {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                "data": base64.b64encode(file_bytes).decode()}}
    if (low.endswith(".heic") or low.endswith(".heif") or "heic" in (mime or "")) and heic_convert:
        file_bytes, mime = heic_convert(file_bytes)
    if downscale:
        # Keep text legible: documents need more pixels than skin photos.
        file_bytes, mime = downscale(file_bytes, mime or "image/jpeg", max_dim=2400, quality=90)
    if mime not in ("image/jpeg", "image/png", "image/webp", "image/gif"):
        mime = "image/jpeg"
    return {"type": "image", "source": {"type": "base64", "media_type": mime,
                                        "data": base64.b64encode(file_bytes).decode()}}


def extract_exams(api_key: str, file_bytes: bytes, mime: str, name: str, lang: str = "el",
                  heic_convert=None, downscale=None, log=None) -> dict:
    """Returns {"ok": bool, "file": name, "exams": [...], "patient_name": str|None, "error": str|None}."""
    t0 = time.time()
    try:
        block = _prep_file(file_bytes, mime, name, heic_convert, downscale)
        reply = _call_claude(api_key, _EXTRACT_SYSTEM, [block, {"type": "text", "text": _extract_prompt(lang)}])
        data = _parse_json(reply)
        exams = [_clean_exam(e) for e in (data.get("exams") or []) if isinstance(e, dict)]
        for e in exams:
            e["files"] = [name]
        if log:
            log("dossier_extract", ok=True, ms=(time.time() - t0) * 1000, n_exams=len(exams))
        return {"ok": True, "file": name, "exams": exams,
                "patient_name": data.get("patient_name_on_document"), "error": None}
    except Exception as ex:  # noqa: BLE001
        if log:
            log("dossier_extract", ok=False, ms=(time.time() - t0) * 1000, error=str(ex)[:200])
        return {"ok": False, "file": name, "exams": [], "patient_name": None, "error": str(ex)}


def extract_many(api_key: str, files: list, lang: str = "el", heic_convert=None, downscale=None,
                 log=None, max_workers: int = 4) -> list:
    """files: list of (bytes, mime, name). Runs extractions in parallel (keeps total wait short,
    which matters on phones that suspend background tabs)."""
    with _cf.ThreadPoolExecutor(max_workers=max_workers) as ex:
        futs = [ex.submit(extract_exams, api_key, b, m, n, lang, heic_convert, downscale, log)
                for (b, m, n) in files]
        return [f.result() for f in futs]


def _s(v):
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def _clean_exam(e: dict) -> dict:
    secs = []
    for sec in e.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        cols = [str(c) for c in (sec.get("columns") or []) if str(c).strip()] or ["Parameter", "Value"]
        rows = []
        for r in sec.get("rows") or []:
            if isinstance(r, (list, tuple)):
                r = [("" if c is None else str(c)) for c in r]
            elif isinstance(r, dict):
                r = [str(r.get(k, "")) for k in r]
            else:
                continue
            r = (r + [""] * len(cols))[:len(cols)]
            if any(x.strip() for x in r):
                rows.append(r)
        if rows or _s(sec.get("description")):
            secs.append({"title": _s(sec.get("title")), "description": _s(sec.get("description")),
                         "columns": cols, "rows": rows})
    concl = []
    for c in e.get("conclusion") or []:
        if isinstance(c, dict) and (_s(c.get("item")) or _s(c.get("detail"))):
            concl.append({"item": _s(c.get("item")) or "", "detail": _s(c.get("detail")) or ""})
        elif isinstance(c, str) and c.strip():
            concl.append({"item": "", "detail": c.strip()})
    return {
        "exam_title": _s(e.get("exam_title")) or "—",
        "exam_type": _s(e.get("exam_type")) or "other",
        "date": _s(e.get("date")), "date_iso": _s(e.get("date_iso")),
        "facility": _s(e.get("facility")), "clinician": _s(e.get("clinician")),
        "source_note": _s(e.get("source_note")),
        "sections": secs, "conclusion": concl,
    }


def merge_exams(exams: list) -> list:
    """Merge pages of the same exam (same title + date) that arrived as separate photos,
    then sort most-recent first (undated last)."""
    merged: dict = {}
    order = []
    for e in exams:
        key = ((e.get("exam_title") or "").strip().lower(), e.get("date_iso") or e.get("date") or "")
        if key in merged:
            m = merged[key]
            m["sections"].extend(e.get("sections") or [])
            m["conclusion"].extend(e.get("conclusion") or [])
            for f in ("facility", "clinician", "source_note"):
                m[f] = m.get(f) or e.get(f)
            m["files"] = list(dict.fromkeys((m.get("files") or []) + (e.get("files") or [])))
        else:
            merged[key] = {**e, "sections": list(e.get("sections") or []),
                           "conclusion": list(e.get("conclusion") or [])}
            order.append(key)

    def _k(e):
        iso = e.get("date_iso")
        if not iso and e.get("date"):
            try:
                iso = datetime.strptime(e["date"].replace(".", "/").replace("-", "/"), "%d/%m/%Y").strftime("%Y-%m-%d")
            except Exception:
                iso = None
        return iso or "0000-00-00"

    return sorted((merged[k] for k in order), key=_k, reverse=True)


def _exam_iso(e: dict):
    iso = e.get("date_iso")
    if not iso and e.get("date"):
        try:
            iso = datetime.strptime(e["date"].replace(".", "/").replace("-", "/"), "%d/%m/%Y").strftime("%Y-%m-%d")
        except Exception:
            iso = None
    return iso


def trend_table(exams: list, max_dates: int = 5, max_rows: int = 40):
    """Parameters that appear in exams on two or more dates, side by side.
    Values are copied as printed — nothing is converted or compared.
    Returns {"columns": [...], "rows": [[param, v_newest, ...], ...]} or None."""
    by_param: dict = {}
    label: dict = {}
    order: list = []
    for e in exams or []:
        iso = _exam_iso(e)
        if not iso:
            continue
        for sec in e.get("sections") or []:
            cols = sec.get("columns") or []
            if len(cols) < 2:
                continue
            for r in sec.get("rows") or []:
                if len(r) < 2 or not str(r[0]).strip() or not str(r[1]).strip():
                    continue
                key = re.sub(r"\s+", " ", str(r[0]).strip().lower())
                if key not in by_param:
                    by_param[key] = {}
                    label[key] = str(r[0]).strip()
                    order.append(key)
                by_param[key].setdefault(iso, str(r[1]).strip())
    keep = [k for k in order if len(by_param[k]) >= 2]
    if not keep:
        return None
    dates = sorted({d for k in keep for d in by_param[k]}, reverse=True)[:max_dates]

    def _fmt(iso):
        try:
            return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception:
            return iso
    rows = [[label[k]] + [by_param[k].get(d, "—") for d in dates] for k in keep[:max_rows]]
    return {"columns": [""] + [_fmt(d) for d in dates], "rows": rows}


def _letterhead_lines(lh: dict):
    lh = lh or {}
    top = _s(lh.get("practice"))
    sub = " · ".join(x for x in [_s(lh.get("doctor")), _s(lh.get("specialty"))] if x)
    contact = " · ".join(x for x in [_s(lh.get("address")), _s(lh.get("phone"))] if x)
    return top, sub, contact


# ── 2. Medication list from a photo (prescription / pill boxes / handwritten) ─
def extract_meds(api_key: str, file_bytes: bytes, mime: str, name: str, lang: str = "el",
                 heic_convert=None, downscale=None, log=None) -> list:
    target = "Greek" if lang == "el" else "English"
    prompt = f"""The attached image/PDF shows a patient's medication (e-prescription, handwritten list,
pill boxes or a hospital discharge note). Return JSON only:
{{"meds": [{{"time": string, "drug": string, "dose": string}}]}}
- "drug": brand name as printed, with the active substance in parentheses if printed or obvious
  from the brand (e.g. "Concor (bisoprolol)").
- "dose": strength and quantity exactly as written (e.g. "1/4 × 10 mg").
- "time": when it is taken if stated (write it in {target}: morning/noon/evening, e.g. "Πρωί"), else "".
- Copy exactly, never invent. Illegible → "[δυσανάγνωστο]" / "[illegible]"."""
    t0 = time.time()
    try:
        block = _prep_file(file_bytes, mime, name, heic_convert, downscale)
        data = _parse_json(_call_claude(api_key, _EXTRACT_SYSTEM,
                                        [block, {"type": "text", "text": prompt}], max_tokens=2500))
        out = []
        for m in data.get("meds") or []:
            if isinstance(m, dict) and _s(m.get("drug")):
                out.append({"time": _s(m.get("time")) or "", "drug": _s(m.get("drug")) or "",
                            "dose": _s(m.get("dose")) or ""})
        if log:
            log("dossier_meds", ok=True, ms=(time.time() - t0) * 1000, n=len(out))
        return out
    except Exception as ex:  # noqa: BLE001
        if log:
            log("dossier_meds", ok=False, ms=(time.time() - t0) * 1000, error=str(ex)[:200])
        return []


# ── 3. Points worth mentioning + symptom summary (text synthesis) ─────────────
def points_worth_mentioning(api_key: str, patient: dict, meds: list, exams: list,
                            lang: str = "el", log=None) -> list:
    """Returns [{"title": str, "text": str}] — factual co-occurrences only."""
    target = "Greek" if lang == "el" else "English"
    system = ("You help a patient prepare for a doctor's appointment by pointing out, for completeness, "
              "facts in their records that sit together or stand out. You are NOT a clinician here: no "
              "diagnosis, no interpretation of meaning, no recommendations, no urgency language.")
    payload = json.dumps({"patient": patient, "medication": meds, "exams": exams}, ensure_ascii=False)
    prompt = f"""Records (JSON):
{payload}

Write 1–5 short "points that may be worth mentioning" in {target}. Allowed kinds of points:
- values the reports THEMSELVES mark as outside the printed normal range, or that exceed a
  normal range printed in the same report;
- related values that appear together across different exams (name each value, exam and date);
- the same parameter measured in more than one exam with different values (list both with dates);
- report-status caveats (e.g. an ECG marked «Unconfirmed Report»);
- a medication and a recorded finding that a doctor would normally want to see side by side
  (state both facts only — no advice, no "should", no "risk").
Each point: a bold-able short title (max 6 words) and 1–3 neutral sentences that end with
wording like "αναφέρεται εδώ για πληρότητα" / "mentioned here for completeness".
Never state what something means, never suggest tests, treatment or urgency.
If nothing qualifies return an empty list.
JSON only: {{"points": [{{"title": string, "text": string}}]}}"""
    t0 = time.time()
    try:
        data = _parse_json(_call_claude(api_key, system, [{"type": "text", "text": prompt}], max_tokens=1800))
        pts = [{"title": _s(p.get("title")) or "", "text": _s(p.get("text")) or ""}
               for p in (data.get("points") or []) if isinstance(p, dict) and _s(p.get("text"))]
        if log:
            log("dossier_points", ok=True, ms=(time.time() - t0) * 1000, n=len(pts))
        return pts[:5]
    except Exception as ex:  # noqa: BLE001
        if log:
            log("dossier_points", ok=False, ms=(time.time() - t0) * 1000, error=str(ex)[:200])
        return []


def summarize_symptoms(api_key: str, conversation: list, lang: str = "el", log=None) -> str:
    """Neutral summary of what the PATIENT said (not what Asklepios suggested)."""
    target = "Greek" if lang == "el" else "English"
    convo = "\n".join(f"{'PATIENT' if m.get('role') == 'user' else 'ASSISTANT'}: {m.get('content', '')[:800]}"
                      for m in conversation or [])
    prompt = f"""Conversation between a patient and a triage assistant:
{convo}

In {target}, write a neutral 3–6 line summary of ONLY what the PATIENT reported: main complaint,
onset/duration, character, associated symptoms, and anything they said they had tried. Use the
patient's own facts; do not include the assistant's assessments, possible causes or advice.
Plain text, short sentences, no heading."""
    t0 = time.time()
    try:
        out = _call_claude(api_key, "You write neutral clinical-history summaries.",
                           [{"type": "text", "text": prompt}], max_tokens=600, timeout=60).strip()
        if log:
            log("dossier_symptoms", ok=True, ms=(time.time() - t0) * 1000)
        return out
    except Exception as ex:  # noqa: BLE001
        if log:
            log("dossier_symptoms", ok=False, ms=(time.time() - t0) * 1000, error=str(ex)[:200])
        return ""


# ── 4. Word document ──────────────────────────────────────────────────────────
def build_docx(dossier: dict, lang: str = "el") -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    lb = labels(lang)
    INK, INDIGO, MUTED = RGBColor(0x0A, 0x10, 0x30), RGBColor(0x43, 0x38, 0xCA), RGBColor(0x5A, 0x63, 0x88)

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def para(text, *, bold=False, italic=False, size=None, color=None, space_after=4, space_before=0):
        p = doc.add_paragraph()
        r = p.add_run(text or "")
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        return p

    def heading(text):
        p = para(text, bold=True, size=14, color=INK, space_before=14, space_after=4)
        return p

    def table(columns, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(columns):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold = True; run.font.size = Pt(10); run.font.color.rgb = INK
            shade(cell, "E0E7FF")
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r[:len(columns)]):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run("" if v is None else str(v))
                run.font.size = Pt(10)
                if i == 0:
                    shade(cells[i], "F4F6FC")
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths[:len(columns)]):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    # Practice letterhead (practice mode)
    _lt, _ls, _lc = _letterhead_lines(dossier.get("letterhead"))
    if _lt or _ls:
        if _lt: para(_lt, bold=True, size=13, color=INDIGO, space_after=0)
        if _ls: para(_ls, size=10, color=INK, space_after=0)
        if _lc: para(_lc, size=9, color=MUTED, space_after=10)

    # Title + patient table
    title = lb["title"]
    if _s(dossier.get("purpose")):
        title = f"{title} — {dossier['purpose']}"
    para(title, bold=True, size=18, color=INK, space_after=8)
    pt = dossier.get("patient") or {}
    prow = []
    if _s(pt.get("name")): prow.append([lb["patient"], pt["name"]])
    if _s(pt.get("amka")): prow.append([lb["amka"], pt["amka"]])
    agesex = ", ".join(x for x in [f"{pt['age']} {lb['years']}" if _s(pt.get("age")) else None, _s(pt.get("sex"))] if x)
    if agesex: prow.append([lb["agesex"], agesex])
    if _s(pt.get("history")): prow.append([lb["history"], pt["history"]])
    if _s(pt.get("allergies")): prow.append(["Αλλεργίες" if lang == "el" else "Allergies", pt["allergies"]])
    if _s(pt.get("dx")): prow.append([lb["dx"], pt["dx"]])
    if prow:
        table([lb["item"], lb["value"]], prow, widths=[4.5, 12.5])
    para(lb["per_exam"], italic=True, color=MUTED)

    # Medication
    meds = [m for m in (dossier.get("meds") or []) if _s(m.get("drug"))]
    if meds:
        heading(lb["meds_h"])
        para(lb["meds_note"], italic=True, color=MUTED)
        table([lb["time"], lb["drug"], lb["dose"]],
              [[m.get("time", ""), m.get("drug", ""), m.get("dose", "")] for m in meds], widths=[3.2, 7.3, 6.5])

    # Values over time
    _tr = trend_table(dossier.get("exams") or []) if dossier.get("trends", True) else None
    if _tr:
        heading(lb["trend_h"])
        para(lb["trend_note"], italic=True, color=MUTED)
        _cols = [lb["param"]] + _tr["columns"][1:]
        _w = 17.0 / len(_cols)
        table(_cols, _tr["rows"], widths=[max(4.5, _w)] + [(17.0 - max(4.5, _w)) / (len(_cols) - 1)] * (len(_cols) - 1))

    # Symptoms (optional)
    if _s(dossier.get("symptoms")):
        heading(lb["symptoms_h"])
        para(lb["symptoms_note"], italic=True, color=MUTED)
        for line in dossier["symptoms"].splitlines():
            if line.strip():
                para(line.strip())

    # Exams
    n = 0
    for e in dossier.get("exams") or []:
        n += 1
        heading(f"{n}. {e.get('exam_title') or '—'}")
        meta = " · ".join(x for x in [e.get("date"), e.get("facility"), e.get("clinician")] if x)
        if e.get("source_note"):
            meta = (meta + " · " if meta else "") + f"{lb['source_note']}: {e['source_note']}"
        if meta:
            para(meta, italic=True, color=MUTED)
        for s in e.get("sections") or []:
            if s.get("title"):
                para(s["title"], bold=True, color=INDIGO, space_before=4)
            if s.get("description"):
                para(f"{s['description']}", italic=True, color=MUTED)
            if s.get("rows"):
                cols = s.get("columns") or [lb["param"], lb["value"]]
                table(cols, s["rows"], widths=[6.5, 10.5] if len(cols) == 2 else [6, 5.5, 5.5])
        if e.get("conclusion"):
            para(lb["conclusion"], bold=True, color=INDIGO, space_before=4)
            if all(not c.get("item") for c in e["conclusion"]):
                for c in e["conclusion"]:
                    para(f"«{c.get('detail', '')}»", italic=True)
            else:
                table([lb["finding"], lb["details"]], [[c.get("item", ""), c.get("detail", "")] for c in e["conclusion"]],
                      widths=[4.5, 12.5])

    # Points worth mentioning
    pts = [p for p in (dossier.get("points") or []) if _s(p.get("text"))]
    if pts:
        heading(f"{n + 1}. {lb['points_h']}")
        para(lb["points_note"], italic=True, color=MUTED)
        for p in pts:
            if _s(p.get("title")):
                para(p["title"], bold=True, space_before=4, space_after=2)
            para(p["text"])

    para(lb["footer"], italic=True, color=MUTED, space_before=14, size=9.5)
    para(f"{lb['made_with']} · {lb['prepared']}: {datetime.now().strftime('%d/%m/%Y')}",
         color=MUTED, size=8.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── 5. Printable HTML ─────────────────────────────────────────────────────────
def build_html(dossier: dict, lang: str = "el") -> str:
    lb = labels(lang)
    esc = lambda x: _html.escape("" if x is None else str(x))  # noqa: E731

    def tbl(cols, rows, first_col_w=None):
        th = "".join(f"<th>{esc(c)}</th>" for c in cols)
        tr = "".join("<tr>" + "".join(f"<td>{esc(v)}</td>" for v in r[:len(cols)]) + "</tr>" for r in rows)
        cg = f'<colgroup><col style="width:{first_col_w}"></colgroup>' if first_col_w else ""
        return f"<table>{cg}<thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>"

    pt = dossier.get("patient") or {}
    prow = []
    if _s(pt.get("name")): prow.append([lb["patient"], pt["name"]])
    if _s(pt.get("amka")): prow.append([lb["amka"], pt["amka"]])
    agesex = ", ".join(x for x in [f"{pt['age']} {lb['years']}" if _s(pt.get("age")) else None, _s(pt.get("sex"))] if x)
    if agesex: prow.append([lb["agesex"], agesex])
    if _s(pt.get("history")): prow.append([lb["history"], pt["history"]])
    if _s(pt.get("allergies")): prow.append(["Αλλεργίες" if lang == "el" else "Allergies", pt["allergies"]])
    if _s(pt.get("dx")): prow.append([lb["dx"], pt["dx"]])

    title = lb["title"] + (f" — {esc(dossier['purpose'])}" if _s(dossier.get("purpose")) else "")
    _lt, _ls, _lc = _letterhead_lines(dossier.get("letterhead"))
    if _lt or _ls:
        _brand = ('<div class="lh">' + (f'<div class="lh-t">{esc(_lt)}</div>' if _lt else "")
                  + (f'<div class="lh-s">{esc(_ls)}</div>' if _ls else "")
                  + (f'<div class="lh-c">{esc(_lc)}</div>' if _lc else "") + '</div>')
    else:
        _brand = f'<div class="brand">⚕ Asklepios · <b>{"AI Νοσηλευτής" if lang == "el" else "AI Nurse"}</b></div>'
    parts = [_brand, f"<h1>{title}</h1>"]
    if prow:
        parts.append(tbl([lb["item"], lb["value"]], prow, "32%"))
    parts.append(f'<p class="note">{esc(lb["per_exam"])}</p>')

    meds = [m for m in (dossier.get("meds") or []) if _s(m.get("drug"))]
    if meds:
        parts.append(f"<h2>{esc(lb['meds_h'])}</h2><p class='note'>{esc(lb['meds_note'])}</p>")
        parts.append(tbl([lb["time"], lb["drug"], lb["dose"]],
                         [[m.get("time", ""), m.get("drug", ""), m.get("dose", "")] for m in meds], "20%"))
    _tr = trend_table(dossier.get("exams") or []) if dossier.get("trends", True) else None
    if _tr:
        parts.append(f"<h2>{esc(lb['trend_h'])}</h2><p class='note'>{esc(lb['trend_note'])}</p>")
        parts.append(tbl([lb["param"]] + _tr["columns"][1:], _tr["rows"], "30%"))
    if _s(dossier.get("symptoms")):
        parts.append(f"<h2>{esc(lb['symptoms_h'])}</h2><p class='note'>{esc(lb['symptoms_note'])}</p>")
        parts.append("".join(f"<p>{esc(l)}</p>" for l in dossier["symptoms"].splitlines() if l.strip()))

    n = 0
    for e in dossier.get("exams") or []:
        n += 1
        parts.append(f'<section class="exam"><h2><span class="num">{n:02d}</span>{esc(e.get("exam_title"))}</h2>')
        meta = " · ".join(esc(x) for x in [e.get("date"), e.get("facility"), e.get("clinician")] if x)
        if e.get("source_note"):
            meta = (meta + " · " if meta else "") + f"{esc(lb['source_note'])}: {esc(e['source_note'])}"
        if meta:
            parts.append(f'<p class="meta">{meta}</p>')
        for s in e.get("sections") or []:
            if s.get("title"):
                parts.append(f"<h3>{esc(s['title'])}</h3>")
            if s.get("description"):
                parts.append(f'<p class="note">{esc(s["description"])}</p>')
            if s.get("rows"):
                cols = s.get("columns") or [lb["param"], lb["value"]]
                parts.append(tbl(cols, s["rows"], "40%" if len(cols) == 2 else "36%"))
        if e.get("conclusion"):
            parts.append(f"<h3>{esc(lb['conclusion'])}</h3>")
            if all(not c.get("item") for c in e["conclusion"]):
                parts.append("".join(f"<p class='quote'>«{esc(c.get('detail'))}»</p>" for c in e["conclusion"]))
            else:
                parts.append(tbl([lb["finding"], lb["details"]],
                                 [[c.get("item", ""), c.get("detail", "")] for c in e["conclusion"]], "26%"))
        parts.append("</section>")

    pts = [p for p in (dossier.get("points") or []) if _s(p.get("text"))]
    if pts:
        parts.append(f'<section class="exam points"><h2><span class="num">{n + 1:02d}</span>{esc(lb["points_h"])}</h2>'
                     f'<p class="note">{esc(lb["points_note"])}</p>')
        for p in pts:
            parts.append(f"<h3>{esc(p.get('title'))}</h3><p>{esc(p.get('text'))}</p>")
        parts.append("</section>")

    parts.append(f'<p class="footer">{esc(lb["footer"])}</p>')
    parts.append(f'<p class="made">{esc(lb["made_with"])} · {esc(lb["prepared"])}: {datetime.now().strftime("%d/%m/%Y")}</p>')

    btn = "Εκτύπωση / Αποθήκευση ως PDF" if lang == "el" else "Print / Save as PDF"
    return f"""<!doctype html><html lang="{lang}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(lb['title'])} — {esc(pt.get('name') or '')}</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Sora:wght@600;700&family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@600&display=swap');
*{{box-sizing:border-box}}
body{{margin:0;background:#F4F6FC;color:#0A1030;font:14px/1.5 Inter,system-ui,sans-serif}}
.page{{max-width:880px;margin:24px auto;background:#fff;border:1px solid #E1E5F4;border-radius:22px;padding:40px 44px}}
.lh{{border-bottom:2px solid #4F46E5;padding-bottom:12px;margin-bottom:18px}}
.lh-t{{font:700 17px Sora,Inter,sans-serif;color:#4338CA}}
.lh-s{{font-size:13.5px;color:#0A1030;margin-top:2px}}
.lh-c{{font-size:12px;color:#5A6388;margin-top:2px}}
.brand{{font:600 12px 'JetBrains Mono',monospace;letter-spacing:.08em;color:#4F46E5;text-transform:uppercase;margin-bottom:10px}}
h1{{font:700 26px/1.15 Sora,Inter,sans-serif;letter-spacing:-.02em;margin:0 0 18px}}
h2{{font:700 18px/1.25 Sora,Inter,sans-serif;letter-spacing:-.01em;margin:30px 0 4px;display:flex;gap:10px;align-items:baseline}}
h3{{font:600 13.5px Inter,sans-serif;color:#4338CA;margin:16px 0 6px}}
.num{{font:600 12px 'JetBrains Mono',monospace;color:#4F46E5}}
.meta{{color:#5A6388;font-style:italic;margin:0 0 8px}}
.note{{color:#5A6388;font-style:italic;margin:4px 0 8px}}
.quote{{font-style:italic}}
table{{width:100%;border-collapse:collapse;margin:6px 0 12px;font-size:13px;page-break-inside:auto}}
tr{{page-break-inside:avoid}}
th{{background:#E0E7FF;text-align:left;font-weight:700;padding:7px 10px;border:1px solid #C7CEF0}}
td{{padding:7px 10px;border:1px solid #E1E5F4;vertical-align:top}}
td:first-child{{background:#F8F9FE}}
.footer{{margin-top:28px;color:#5A6388;font-style:italic;font-size:12.5px;border-top:1px solid #E1E5F4;padding-top:14px}}
.made{{color:#98A2C8;font-size:11px}}
.bar{{max-width:880px;margin:18px auto 0;text-align:right;padding:0 8px}}
.bar button{{border:0;border-radius:999px;padding:12px 22px;font:700 14px Inter,sans-serif;color:#050816;
  background:linear-gradient(135deg,#818CF8,#22D3EE);cursor:pointer}}
@media (max-width:640px){{.page{{margin:10px;padding:22px 16px;border-radius:16px}} h1{{font-size:21px}} table{{font-size:12px}}}}
@media print{{body{{background:#fff}} .bar{{display:none}} .page{{border:0;margin:0;padding:0;max-width:none}}
  th{{-webkit-print-color-adjust:exact;print-color-adjust:exact}} td:first-child{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
  .exam{{page-break-inside:auto}} h2,h3{{page-break-after:avoid}}}}
</style></head><body>
<div class="bar"><button onclick="window.print()">🖨 {btn}</button></div>
<div class="page">{''.join(parts)}</div>
</body></html>"""
